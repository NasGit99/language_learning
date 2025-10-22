import os
import sys
import pytest
from docx import Document

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '../')))

from src.language_modifier.docx_translator import DocxTranslator

test_dir = os.path.join(os.path.dirname(__file__), "") 
os.makedirs(test_dir, exist_ok=True)

test_file_1_name ="test_file_1.docx"
test_file_1_path = os.path.join(test_dir, test_file_1_name)


@pytest.fixture(scope="session",autouse=True)
def create_docx_file():
    
    doc = Document()
    doc.add_heading("Test Document", level=1)
    doc.add_paragraph("This is a sample paragraph inside the test document.")
    para = doc.add_paragraph("This text is in bold and italic: ")
    para.add_run("Bold ").bold = True
    para.add_run("Italic").italic = True
    doc.add_paragraph("First bullet point", style="List Bullet")
    doc.add_paragraph("Second bullet point", style="List Bullet")
    doc.add_paragraph("First item", style="List Number")
    doc.add_paragraph("Second item", style="List Number")
    doc.save(test_file_1_name)

class TestInput():
    def test_docx_translation(self):
        file_1 = DocxTranslator(test_file_1_name, "Spanish", test_dir,testing=True )
        output = file_1.generate_doc()
        assert os.path.exists(os.path.join(test_dir, output))
    
    def test_json_docx_file(self, client,json_users):
        import io
        user_1, _ = json_users
        file_data = {
            "file": (io.BytesIO(b"dummy content"), test_file_1_name),
            "target_language": "French"
        }

        response_send_file = client.post(
            "/translate_document",
            data=file_data,
            content_type="multipart/form-data",
            headers={"Authorization": f"Bearer {user_1.token}"}
        )

        assert response_send_file.status_code == 200

