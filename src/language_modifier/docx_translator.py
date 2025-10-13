from language_modifier.base_translator import TranslatorCore
from language_modifier.language_translator import translate_text
import asyncio
from docx import Document


class DocxTranslator(TranslatorCore):
    def __init__(self, file_path, target_lang_code, upload_folder=None, username=None, testing=None):
        super().__init__(file_path, target_lang_code, upload_folder,testing)
        self.username = username
    
    # Read the file
def doc_reader():
    doc = Document("sample3.docx")
    # This can probably be cleaned up a bit for separate functions
    rows_str = ""

    for table in doc.tables:
        rows = table.rows

        for row in rows:
            for cell in row.cells:
                rows_str += cell.text + "|||"

    paragraph_str = ""

    for paragraph in doc.paragraphs:
        paragraph_str += paragraph.text + "|||" 
    
    header_tbl_str = ""
    footer_tbl_str = ""
    header_text_str = ""
    footer_text_str = ""

    for section in doc.sections:
        # We can just define section.header section.footer than do the loops
        header_tbl= section.header.tables
        footer_tbl = section.footer.tables
        header_paragraphs = section.header.paragraphs
        footer_paragraphs = section.footer.paragraphs

        if header_tbl:
            for table in header_tbl:
                rows = table.rows
                for row in rows:
                    for cell in row.cells:
                        header_tbl_str += cell.text + "|||"

        if footer_tbl:
            for table in header_tbl:
                rows = table.rows
                for row in rows:
                    for cell in row.cells:
                        footer_tbl_str + "|||"
        
        if header_paragraphs:
            for paragraph in header_paragraphs:
                header_text_str += paragraph.text + "|||"
        
        if footer_paragraphs:
            for paragraph in footer_paragraphs:
                footer_text_str += paragraph.text + "|||"
                     
    return doc, rows_str, paragraph_str, header_tbl_str, footer_tbl_str, header_text_str, footer_text_str

def translate_strings(text):
    translated_text = asyncio.run(translate_text(text,"spanish"))
    return translated_text

def translate_doc():
    doc, rows_str, paragraph_str, header_tbl_str, footer_tbl_str, \
    header_text_str, footer_text_str = doc_reader()

    translated_rows = translate_strings(rows_str)
    translated_paragraphs = translate_strings(paragraph_str)
    translated_header_tbl = translate_strings(header_tbl_str)
    translated_footer_tbl = translate_strings(footer_tbl_str)
    translated_header_text = translate_strings(header_text_str)
    translated_footer_text = translate_strings(footer_text_str)

    return doc, translated_rows, translated_paragraphs, translated_header_tbl, translated_footer_tbl, \
    translated_header_text, translated_footer_text

def generate_doc():
    doc, translated_rows, translated_paragraphs, translated_header_tbl, translated_footer_tbl, \
    translated_header_text, translated_footer_text = translate_doc()

    translated_rows.split("|||")
    print(translated_rows)

    # for table in doc.tables:
    #     rows = table.rows

    #     for row in rows:
    #         for cell in row.cells:
    #             for word in translated_rows:
    #                 cell.text = word
    #                 print(word)






generate_doc()